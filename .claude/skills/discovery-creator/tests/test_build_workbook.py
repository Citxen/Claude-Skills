import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))

try:
    import build_workbook
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except SystemExit:  # build_workbook exits when openpyxl is missing
    HAS_OPENPYXL = False
except ImportError:
    HAS_OPENPYXL = False

SPEC = {
    "sheets": [
        {"name": "Discovery register", "title": "Discovery register", "subtitle": "One row per question.",
         "headers": ["ID", "Domain", "Flag", "Answer", "Status"],
         "widths": [8, 20, 30, 20, 10],
         "rows": [["Q001", "Scope", "", "", "Open"],
                  ["Q002", "Scope", "GOV GAP: not available in DoD", "", "Open"],
                  ["EX-1", "Scope", "", "", "Open"]],
         "example_rows": [3],
         "input_columns": ["D", "E"],
         "dropdowns": {"E": ["Open", "Answered"]},
         "flag_column": 3,
         "filter": True},
        {"name": "Dashboard", "headers": ["Domain", "Questions"],
         "rows": [["Scope", "=COUNTIF('Discovery register'!$B$5:$B$7,$A5)"]]},
    ]
}

@unittest.skipUnless(HAS_OPENPYXL, "openpyxl is not installed")
class BuildWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / "pack.xlsx"
        build_workbook.build(SPEC, self.path)
        self.workbook = load_workbook(self.path)

    def tearDown(self):
        self.workbook.close()
        self.directory.cleanup()

    def test_sheets_in_spec_order(self):
        self.assertEqual(self.workbook.sheetnames, ["Discovery register", "Dashboard"])

    def test_title_subtitle_and_header_row(self):
        sheet = self.workbook["Discovery register"]
        self.assertEqual(sheet["A1"].value, "Discovery register")
        self.assertEqual(sheet["A2"].value, "One row per question.")
        self.assertEqual([cell.value for cell in sheet[4]], ["ID", "Domain", "Flag", "Answer", "Status"])

    def test_rows_start_below_the_header(self):
        sheet = self.workbook["Discovery register"]
        self.assertEqual(sheet["A5"].value, "Q001")
        self.assertEqual(sheet["A7"].value, "EX-1")

    def test_flagged_row_is_shaded_and_unflagged_row_is_not(self):
        sheet = self.workbook["Discovery register"]
        self.assertEqual(sheet["A6"].fill.fgColor.rgb[-6:], "FCE4D6")
        self.assertNotEqual(sheet["A5"].fill.fgColor.rgb[-6:], "FCE4D6")

    def test_input_columns_are_shaded(self):
        self.assertEqual(self.workbook["Discovery register"]["D5"].fill.fgColor.rgb[-6:], "FFF2CC")

    def test_example_row_is_greyed(self):
        self.assertEqual(self.workbook["Discovery register"]["A7"].font.color.rgb[-6:], "808080")

    def test_dropdown_is_attached(self):
        sheet = self.workbook["Discovery register"]
        self.assertTrue(any("Open,Answered" in (validation.formula1 or "")
                            for validation in sheet.data_validations.dataValidation))

    def test_formula_is_written_as_a_formula(self):
        self.assertTrue(str(self.workbook["Dashboard"]["B5"].value).startswith("=COUNTIF("))

    def test_filter_and_freeze(self):
        sheet = self.workbook["Discovery register"]
        self.assertEqual(sheet.auto_filter.ref, "A4:E7")
        self.assertEqual(sheet.freeze_panes, "A5")

    def test_font_is_arial_throughout(self):
        sheet = self.workbook["Discovery register"]
        self.assertEqual(sheet["A5"].font.name, "Arial")
        self.assertEqual(sheet["A4"].font.name, "Arial")

@unittest.skipUnless(HAS_OPENPYXL, "openpyxl is not installed")
class CommandLineTest(unittest.TestCase):
    def run_script(self, *args):
        return subprocess.run([sys.executable, str(SCRIPTS / "build_workbook.py"), *args],
                              capture_output=True, text=True)

    def test_usage_error(self):
        self.assertEqual(self.run_script().returncode, 2)

    def test_unreadable_spec(self):
        self.assertEqual(self.run_script("does-not-exist.json").returncode, 2)

if __name__ == '__main__':
    unittest.main()
