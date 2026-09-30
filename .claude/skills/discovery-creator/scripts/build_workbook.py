"""Build a discovery workbook from a JSON specification.

The spec carries the content; this script owns the formatting, so every pack this skill
produces looks the same. See references/workbook_spec.md for the shape of the spec.
"""
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, column_index_from_string
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - exercised only where the dependency is absent
    print("error: openpyxl is required to build a workbook: python -m pip install openpyxl",
          file=sys.stderr)
    raise SystemExit(2)

FONT = "Arial"
HEADER_ROW = 4
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, size=14, bold=True, color="1F3864")
SUBTITLE_FONT = Font(name=FONT, size=10, italic=True, color="595959")
BODY_FONT = Font(name=FONT, size=10)
EXAMPLE_FONT = Font(name=FONT, size=10, italic=True, color="808080")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")
EDGE = Side(style="thin", color="BFBFBF")
BORDER = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)
WRAP = Alignment(wrap_text=True, vertical="top")

def add_sheet(workbook, spec):
    sheet = workbook.create_sheet(spec["name"])
    sheet["A1"] = spec.get("title", spec["name"])
    sheet["A1"].font = TITLE_FONT
    if spec.get("subtitle"):
        sheet["A2"] = spec["subtitle"]
        sheet["A2"].font = SUBTITLE_FONT

    headers = spec["headers"]
    widths = spec.get("widths") or [20] * len(headers)
    for number, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(row=HEADER_ROW, column=number, value=header)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, WRAP, BORDER
        sheet.column_dimensions[get_column_letter(number)].width = width
    sheet.row_dimensions[HEADER_ROW].height = 30

    examples = set(spec.get("example_rows") or [])
    flag_column = spec.get("flag_column")
    flag_prefix = spec.get("flag_prefix", "GOV GAP")
    rows = spec.get("rows") or []
    for offset, row in enumerate(rows, start=1):
        is_example = offset in examples
        flagged = bool(flag_column) and str(row[flag_column - 1] if len(row) >= flag_column else "").startswith(flag_prefix)
        for number, value in enumerate(row, start=1):
            cell = sheet.cell(row=HEADER_ROW + offset, column=number, value=value)
            cell.font = EXAMPLE_FONT if is_example else BODY_FONT
            cell.alignment, cell.border = WRAP, BORDER
            if flagged:
                cell.fill = FLAG_FILL

    last = HEADER_ROW + len(rows)
    for letter in spec.get("input_columns") or []:
        for number in range(HEADER_ROW + 1, last + 1):
            if not (spec.get("flag_column") and sheet[f"{letter}{number}"].fill == FLAG_FILL):
                sheet[f"{letter}{number}"].fill = INPUT_FILL

    for letter, options in (spec.get("dropdowns") or {}).items():
        validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"{letter}{HEADER_ROW + 1}:{letter}{max(last, HEADER_ROW + 1)}")

    if spec.get("filter") and rows:
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(headers))}{last}"
    sheet.freeze_panes = spec.get("freeze", f"A{HEADER_ROW + 1}")
    return sheet

def build(spec, output=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_spec in spec["sheets"]:
        add_sheet(workbook, sheet_spec)
    destination = Path(output or spec["output"])
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination

if __name__ == '__main__':
    if not 2 <= len(sys.argv) <= 3:
        print("usage: build_workbook.py <spec.json> [output.xlsx]", file=sys.stderr)
        sys.exit(2)
    try:
        specification = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as error:
        print(f"error: cannot read spec: {error}", file=sys.stderr)
        sys.exit(2)
    try:
        path = build(specification, sys.argv[2] if len(sys.argv) == 3 else None)
    except (OSError, KeyError, TypeError) as error:
        print(f"error: cannot build workbook: {error}", file=sys.stderr)
        sys.exit(1)
    sheets = len(specification["sheets"])
    rows = sum(len(sheet.get("rows") or []) for sheet in specification["sheets"])
    print(f"wrote {path}: {sheets} sheets, {rows} rows")
